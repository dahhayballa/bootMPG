"""
Commande d'import automatique des documents .docx réels (Étape 4).

Remplace la saisie manuelle via l'admin pour les documents volumineux : on
extrait le texte (paragraphes ET tableaux, dans l'ordre réel du document),
on découpe par section (Heading 2 = 1 chunk), et on catégorise chaque
section via son numéro (ex: "3. Admission et inscription" -> categorie
'admission'), en s'appuyant sur le fait que la numérotation est identique
entre les versions française et arabe de chaque document.

Fichiers attendus dans knowledge/seed_documents/ :
    01_Guide_Etudiant_AR.docx      02_Guide_Etudiant_FR.docx
    03_Programmes_et_Filieres_FR.docx   04_Programmes_et_Filieres_AR.docx
    05_FAQ_Bot_AR.docx             06_FAQ_Bot_FR.docx

RÈGLE DE PRUDENCE APPLIQUÉE ICI (à valider avec l'utilisateur, cf. README) :
- Guide de l'étudiant et Programmes/Filières -> chunks importés en statut
  'en_attente_validation' (contenu narratif dense, un passage humain reste
  recommandé avant activation - conforme à la règle Étape 4).
- FAQ (Bot) -> importées directement en statut 'actif', car ce fichier est
  explicitement conçu et rédigé comme base de connaissances prête pour le
  bot (format 1 question = 1 réponse autonome, marquage ✅/⚠️/🕐 déjà
  appliqué par la source). Cette exception doit être confirmée par
  l'utilisateur (cf. réponse accompagnant cette commande).

Usage :
    python manage.py import_docx_documents
"""

import re
from datetime import date
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from django.core.management.base import BaseCommand
from django.db import transaction

from knowledge.models import FAQ, Chunk, Document

SEED_DIR = Path(__file__).resolve().parent.parent.parent / 'seed_documents'

DATE_VERSION = date(2026, 9, 1)  # "Version 1.0 — septembre 2026", indiquée dans chaque document

HEADING_NUM_RE = re.compile(r'^(\d+)\.\s*')
QA_PREFIX_RE = re.compile(r'^(?:Q\s*:|س\s*:)\s*')
SYMBOL_RE = re.compile(r'(✅|⚠️|🕐)')


def nettoyer_reponse_faq(reponse: str) -> str:
    """Retire les consignes éditoriales qui ne sont pas destinées au public."""
    reponse = reponse.strip()
    if reponse.startswith('الجواب الصحيح للبوت:'):
        reponse = reponse.removeprefix('الجواب الصحيح للبوت:').strip(' «»"')
    elif reponse.startswith('Réponse correcte pour le bot :'):
        reponse = reponse.removeprefix('Réponse correcte pour le bot :').strip(' «»"')
    reponse = reponse.replace('المعلومات الموسومة 🕐 مقترحة', 'المعلومات المقترحة')
    reponse = reponse.replace('الموسومة ⚠️ غير متوفرة', 'غير المتوفرة')
    return reponse.replace('⚠️', 'ملاحظة:').replace('🕐', 'ملاحظة:').strip()

# Sections narratives (Guide, Programmes) : numéro de Heading 2 -> catégorie.
# La numérotation est identique en français et en arabe (vérifié sur les 6 fichiers).
GUIDE_SECTION_CATEGORIE = {
    1: 'ecole', 2: 'specialites', 3: 'admission', 4: 'etudes', 5: 'etudes',
    6: 'examens', 7: 'stages', 8: 'reglement', 9: 'reglement',
    10: 'services_etudiants', 11: 'services_etudiants', 12: 'etudes', 13: 'contacts',
}
PROGRAMMES_SECTION_CATEGORIE = {n: 'specialites' for n in range(0, 13)}
FAQ_SECTION_CATEGORIE = {
    1: 'ecole', 2: 'specialites', 3: 'admission', 4: 'etudes', 5: 'etudes',
    6: 'examens', 7: 'stages', 8: 'reglement', 9: 'reglement',
    10: 'services_etudiants', 11: 'services_etudiants', 12: 'faq',
}
# Section 0 du Guide = pur mode d'emploi pour le bot, sans contenu factuel -> exclue.
GUIDE_SECTION_CATEGORIE.setdefault(0, None)


def iter_block_items(document):
    """Parcourt le corps du document dans l'ordre réel (paragraphes ET
    tableaux mélangés), contrairement à `document.paragraphs` qui ignore
    les tableaux et `document.tables` qui perd leur position. Recette
    standard python-docx.
    """
    parent_elm = document.element.body
    for child in parent_elm.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, document)
        elif child.tag == qn('w:tbl'):
            yield Table(child, document)


def table_to_text(table: Table) -> str:
    lignes = []
    for row in table.rows:
        cellules = [c.text.strip() for c in row.cells]
        if any(cellules):
            lignes.append(" | ".join(cellules))
    return "\n".join(lignes)


def extraire_sections(docx_path: Path, section_categorie_map: dict) -> list[dict]:
    """Découpe un document en sections (1 par Heading 2), en conservant
    l'ordre naturel du texte et des tableaux. Retourne une liste de dicts
    {numero, titre, categorie, texte}.
    """
    document = docx.Document(str(docx_path))
    sections = []
    section_courante = None

    def flush():
        if section_courante and section_courante['categorie'] and section_courante['blocs']:
            texte = section_courante['titre'] + "\n\n" + "\n\n".join(section_courante['blocs'])
            sections.append({
                'numero': section_courante['numero'],
                'titre': section_courante['titre'],
                'categorie': section_courante['categorie'],
                'texte': texte.strip(),
            })

    for block in iter_block_items(document):
        if isinstance(block, Paragraph):
            style = block.style.name if block.style else ''
            texte = block.text.strip()
            if not texte:
                continue
            if style == 'Heading 1':
                continue  # titre du document, déjà connu
            if style == 'Heading 2':
                flush()
                m = HEADING_NUM_RE.match(texte)
                numero = int(m.group(1)) if m else None
                section_courante = {
                    'numero': numero,
                    'titre': texte,
                    'categorie': section_categorie_map.get(numero) if numero is not None else None,
                    'blocs': [],
                }
                continue
            if section_courante is None:
                continue
            if style == 'Heading 3':
                section_courante['blocs'].append(f"— {texte} —")
            else:
                section_courante['blocs'].append(texte)
        elif isinstance(block, Table):
            if section_courante is None:
                continue
            texte_table = table_to_text(block)
            if texte_table:
                section_courante['blocs'].append(texte_table)

    flush()
    return sections


def extraire_faq(docx_path: Path, section_categorie_map: dict) -> list[dict]:
    """Extrait les paires question/réponse d'un fichier FAQ, section par
    section, en s'appuyant sur le format "Q : ... <symbole> réponse" déjà
    utilisé par la source (1 question = 1 paragraphe autonome).
    """
    document = docx.Document(str(docx_path))
    resultats = []
    categorie_courante = None

    for para in document.paragraphs:
        style = para.style.name if para.style else ''
        texte = para.text.strip()
        if not texte:
            continue
        if style == 'Heading 2':
            m = HEADING_NUM_RE.match(texte)
            numero = int(m.group(1)) if m else None
            categorie_courante = section_categorie_map.get(numero) if numero is not None else None
            continue
        if categorie_courante is None:
            continue
        if not QA_PREFIX_RE.match(texte):
            continue  # ligne d'intro de section, pas une question
        reste = QA_PREFIX_RE.sub('', texte, count=1)
        m = SYMBOL_RE.search(reste)
        if not m:
            continue
        question = reste[:m.start()].strip(' ?؟') + ('؟' if 'س' in texte[:3] else ' ?')
        reponse = nettoyer_reponse_faq(SYMBOL_RE.sub('', reste[m.start():], count=1))
        resultats.append({'question': question, 'reponse': reponse, 'categorie': categorie_courante})

    return resultats


class Command(BaseCommand):
    help = "Importe les 6 documents .docx réels (Guide, Programmes, FAQ - FR/AR) dans la base de connaissances."

    @transaction.atomic
    def handle(self, *args, **options):
        self._importer_narratif(
            'Guide de l\'étudiant (FR)', '02_Guide_Etudiant_FR.docx', 'fr',
            'guide_etudiant', GUIDE_SECTION_CATEGORIE,
        )
        self._importer_narratif(
            'Guide de l\'étudiant (AR)', '01_Guide_Etudiant_AR.docx', 'ar',
            'guide_etudiant', GUIDE_SECTION_CATEGORIE,
        )
        self._importer_narratif(
            'Programmes de formation et filières (FR)', '03_Programmes_et_Filieres_FR.docx', 'fr',
            'programme_filiere', PROGRAMMES_SECTION_CATEGORIE,
        )
        self._importer_narratif(
            'Programmes de formation et filières (AR)', '04_Programmes_et_Filieres_AR.docx', 'ar',
            'programme_filiere', PROGRAMMES_SECTION_CATEGORIE,
        )
        self._importer_faq('FAQ Bot (FR)', '06_FAQ_Bot_FR.docx', 'fr', FAQ_SECTION_CATEGORIE)
        self._importer_faq('FAQ Bot (AR)', '05_FAQ_Bot_AR.docx', 'ar', FAQ_SECTION_CATEGORIE)

    def _importer_narratif(self, titre_doc, filename, langue, type_document, section_map):
        chemin = SEED_DIR / filename
        if not chemin.exists():
            self.stderr.write(self.style.ERROR(f"Fichier introuvable : {chemin}"))
            return

        sections = extraire_sections(chemin, section_map)

        document, _ = Document.objects.update_or_create(
            titre=titre_doc,
            defaults=dict(
                type_document=type_document,
                fichier_original=filename,
                langue=langue,
                categorie_principale=sections[0]['categorie'] if sections else 'ecole',
                date_creation_document=DATE_VERSION,
                date_derniere_verification=DATE_VERSION,
                statut_officialite='officiel_valide',
                niveau_confiance='haute',
                statut_activite='actif',
                version='1.0',
                notes_internes=(
                    "Importé automatiquement (import_docx_documents). Chunks en "
                    "'en_attente_validation' : à relire puis activer via l'admin."
                ),
            ),
        )
        document.chunks.all().delete()

        for s in sections:
            Chunk.objects.create(
                document=document,
                contenu_texte=s['texte'],
                langue=langue,
                categorie=s['categorie'],
                ordre_dans_document=s['numero'] or 0,
                statut_activite='en_attente_validation',
            )

        self.stdout.write(self.style.SUCCESS(
            f"'{titre_doc}' : {len(sections)} chunks créés (en_attente_validation)."
        ))

    def _importer_faq(self, titre_doc, filename, langue, section_map):
        chemin = SEED_DIR / filename
        if not chemin.exists():
            self.stderr.write(self.style.ERROR(f"Fichier introuvable : {chemin}"))
            return

        paires = extraire_faq(chemin, section_map)

        document, _ = Document.objects.update_or_create(
            titre=titre_doc,
            defaults=dict(
                type_document='faq',
                fichier_original=filename,
                langue=langue,
                categorie_principale='faq',
                date_creation_document=DATE_VERSION,
                date_derniere_verification=DATE_VERSION,
                statut_officialite='officiel_valide',
                niveau_confiance='haute',
                statut_activite='actif',
                version='1.0',
                notes_internes=(
                    "Importé automatiquement (import_docx_documents). Contrairement aux "
                    "chunks narratifs, ces FAQ sont importées directement en statut "
                    "'actif' : ce fichier source est explicitement rédigé comme base "
                    "de connaissances prête pour le bot (1 question = 1 réponse "
                    "autonome, déjà marquée ✅/⚠️/🕐 par l'auteur)."
                ),
            ),
        )
        FAQ.objects.filter(source_document=document).delete()

        for p in paires:
            FAQ.objects.create(
                question=p['question'],
                reponse=p['reponse'],
                langue=langue,
                categorie=p['categorie'],
                source_document=document,
                statut_activite='actif',
            )

        self.stdout.write(self.style.SUCCESS(
            f"'{titre_doc}' : {len(paires)} FAQ créées (actif)."
        ))
